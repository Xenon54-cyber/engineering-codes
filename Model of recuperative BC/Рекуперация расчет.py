import numpy as np
import pandas as pd
from CoolProp.CoolProp import PropsSI as prop
from scipy.optimize import root
from tqdm import tqdm
import itertools
import time
from datetime import datetime
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
import multiprocessing as mp

# ============================================================
# ИСХОДНЫЕ ДАННЫЕ
# ============================================================

FLUID = "CO2"

TR1_C = 730.0
TR2_C = 550.0
T2_SET_C = 715.0
T1_C = 712.0
T4_C = 35.0

TR1 = TR1_C + 273.15
TR2 = TR2_C + 273.15
T2_SET = T2_SET_C + 273.15
T1 = T1_C + 273.15
T4 = T4_C + 273.15
Q_HEAT = 2.750e6
KPDcomp = 0.80
KPDturb = 0.90
KPD_reactor = 0.99
KPD_RTO = 0.99
ETA_T = KPDturb
ETA_C = KPDcomp
ETA_GENERATOR = 0.98
eta_turb_mech = 0.99
eta_generator = 0.99
eta_mc_mech = 0.97
eta_rc_mech = 0.97
eta_drive = 0.98
eta_tr1 = 0.99
eta_tr2 = 0.99
DT_REC = 5.0
Q_reactor = Q_HEAT * KPD_reactor
Q_tr1 = Q_reactor * eta_tr1
Q_CO2_RTO = Q_tr1 * KPD_RTO

Q_loss_reactor = Q_HEAT - Q_reactor
Q_loss_tr1 = Q_reactor - Q_tr1
Q_loss_RTO = Q_tr1 - Q_CO2_RTO

P1_RANGE = np.arange(20.0, 36.1, 1.5) * 1e6
P2_RANGE = np.arange(7.5, 16.1, 1.5) * 1e6
DP_REC_HOT_FRAC = 0.03
DP_REC_COLD_FRAC = 0.02
CP_NA = 1270.0
RHO_NA = 850.0
G_NA = Q_reactor / (CP_NA * (TR1_C - TR2_C))
RHO_WATER = 1000.0
CP_WATER = 4180.0
P_NA = 1.0e6
DP_NA_FRAC = 0.05
P_WATER = 0.50e6
DP_WATER_FRAC = 0.10

ETA_PUMP = 0.80
ETA_MOTOR = 0.95
DT_WATER = 10.0


def H_TP(T, P):
    return prop("H", "T", T, "P", P, FLUID)
def S_TP(T, P):
    return prop("S", "T", T, "P", P, FLUID)
def H_PS(P, S):
    return prop("H", "P", P, "S", S, FLUID)
def T_PH(P, H):
    return prop("T", "P", P, "H", H, FLUID)
def calculate_pressures(P1, P2):
    if P2 >= P1:
        raise ValueError("P2 >= P1")
    P3 = P2 * (1.0 - DP_REC_HOT_FRAC)
    P4 = P3
    P6 = P1
    P5 = P6 / (1.0 - DP_REC_COLD_FRAC)
    return {"P1": P1, "P2": P2, "P3": P3, "P4": P4, "P5": P5, "P6": P6}

def recuperator_profile(P2, P3, P5, P6, h2, h5, T6, n_sections=20):
    h6 = H_TP(T6, P6)
    dh_cold = h6 - h5
    if dh_cold <= 0:
        raise ValueError("h6-h5 <= 0")
    h3 = h2 - dh_cold
    T3 = T_PH(P3, h3)
    frac = np.linspace(0.0, 1.0, n_sections + 1)
    H_hot = h2 - frac * (h2 - h3)
    H_cold_counter = h6 - frac * (h6 - h5)
    T_hot = np.array([T_PH(P2, h) for h in H_hot])
    T_cold = np.array([T_PH(P6, h) for h in H_cold_counter])
    dT_profile = T_hot - T_cold
    dTmin = float(np.min(dT_profile))
    idx_min = int(np.argmin(dT_profile))
    return {"h6": h6,"h3": h3,"T3": T3,"Q_rec_specific": dh_cold,"T_hot_profile": T_hot.tolist(),"T_cold_profile": T_cold.tolist(),"dT_profile": dT_profile.tolist(),"dTmin": dTmin,"idx_min": idx_min}

def calc(args):
    P1_calc, P2_calc = args
    r = {"P1_MPa": P1_calc / 1e6,"P2_MPa": P2_calc / 1e6,"eta_T_%": ETA_T * 100.0,"eta_C_%": ETA_C * 100.0,"physical_root": False,"success": False,"message": ""}
    try:
        pressures = calculate_pressures(P1_calc, P2_calc)
        P1 = pressures["P1"]
        P2 = pressures["P2"]
        P3 = pressures["P3"]
        P4 = pressures["P4"]
        P5 = pressures["P5"]
        P6 = pressures["P6"]
        h1 = H_TP(T1, P1)
        s1 = S_TP(T1, P1)
        h2s = H_PS(P2, s1)
        h2 = h1 - ETA_T * (h1 - h2s)
        T2 = T_PH(P2, h2)
        w_turb = h1 - h2
        h4 = H_TP(T4, P4)
        s4 = S_TP(T4, P4)
        h5s = H_PS(P5, s4)
        h5 = h4 + (h5s - h4) / ETA_C
        T5 = T_PH(P5, h5)
        w_comp = h5 - h4
        h_before_transport = H_TP(T2_SET, P1)
        h_after_transport = H_TP(T1, P1)

        def equations(x):
            G = float(x[0])
            T6 = float(x[1])
            if G <= 0:
                return np.array([1e6, 1e6])
            try:
                h6 = H_TP(T6, P6)
                eq1 = G * (h_before_transport - h6) - Q_CO2_RTO
                rec = recuperator_profile(P2, P3, P5, P6, h2, h5, T6, n_sections=20)
                eq2 = rec["dTmin"] - DT_REC
                return np.array([eq1 / 1e6, eq2])
            except Exception:
                return np.array([1e6, 1e6])
        G0 = Q_CO2_RTO / (h_before_transport - H_TP(TR2 - 5.0, P6))
        if not np.isfinite(G0) or G0 <= 0:
            G0 = 20.0
        x0 = np.array([G0, TR2 - 5.0])
        sol = root(equations, x0=x0, method="hybr", tol=1e-8)
        if not sol.success:
            raise ValueError("ROOT не сошелся: " + str(sol.message))
        G_CO2 = float(sol.x[0])
        T6 = float(sol.x[1])
        residual = equations([G_CO2, T6])
        root_residual_max = float(np.max(np.abs(residual)))

        if root_residual_max > 1e-5:
            raise ValueError(f"ROOT дал слишком большую невязку: {root_residual_max:.3e}")
        dT_reactor_cold = TR2 - T6
        if dT_reactor_cold <= 0:
            raise ValueError(f"TR2 > T6 не выполняется: TR2={TR2_C:.3f} C, T6={T6-273.15:.3f} C")
        rec = recuperator_profile(P2, P3, P5, P6, h2, h5, T6, n_sections=20)
        h6 = rec["h6"]
        h3 = rec["h3"]
        T3 = rec["T3"]

        Q_RECUP = G_CO2 * rec["Q_rec_specific"]
        Q_COOL = G_CO2 * (h3 - h4)
        Q_HEAT_CALC = G_CO2 * (h_before_transport - h6)
        Q_transport_loss_CO2 = G_CO2 * (h_before_transport - h_after_transport)
        Q_CO2_to_turbine = Q_CO2_RTO - Q_transport_loss_CO2
        N_TURB_i = G_CO2 * w_turb
        N_MCOMP_i = G_CO2 * w_comp
        N_RCOMP_i = 0.0
        N_NET_i = N_TURB_i - N_MCOMP_i - N_RCOMP_i
        energy_balance = Q_CO2_to_turbine - Q_COOL - N_NET_i

        DP_NA = P_NA * DP_NA_FRAC
        N_NA_PUMP = G_NA / RHO_NA * DP_NA / ETA_PUMP / ETA_MOTOR
        G_WATER = Q_COOL / (CP_WATER * DT_WATER)
        DP_WATER = P_WATER * DP_WATER_FRAC
        N_WATER_PUMP = G_WATER / RHO_WATER * DP_WATER / ETA_PUMP / ETA_MOTOR
        N_SN = N_NA_PUMP + N_WATER_PUMP
        N_TURB_e = N_TURB_i * eta_turb_mech
        N_TURB_el = N_TURB_e * eta_generator
        N_MCOMP_e = N_MCOMP_i / eta_mc_mech
        N_MCOMP_el = N_MCOMP_e / eta_drive
        N_RCOMP_e = 0.0
        N_RCOMP_el = 0.0
        N_EL_GROSS = N_TURB_el - N_MCOMP_el - N_RCOMP_el
        N_EL_NET = N_EL_GROSS - N_SN
        N_loss_turb_mech = N_TURB_i - N_TURB_e
        N_loss_generator = N_TURB_e - N_TURB_el
        N_loss_turb_em = N_loss_turb_mech + N_loss_generator
        N_loss_mc_mech = N_MCOMP_e - N_MCOMP_i
        N_loss_mc_drive = N_MCOMP_el - N_MCOMP_e
        N_loss_rc_mech = 0.0
        N_loss_rc_drive = 0.0

        N_loss_comp_mech = N_loss_mc_mech + N_loss_rc_mech
        N_loss_comp_drive = N_loss_mc_drive + N_loss_rc_drive
        N_loss_comp_em = N_loss_comp_mech + N_loss_comp_drive
        N_loss_pumps = N_SN
        N_loss_em_total = N_loss_turb_em + N_loss_comp_em + N_loss_pumps

        eta_cycle_CO2 = N_NET_i / Q_CO2_RTO * 100.0
        eta_el_net_CO2 = N_EL_NET / Q_CO2_RTO * 100.0
        eta_plant = N_NET_i / Q_HEAT * 100.0
        eta_el_net_plant = N_EL_NET / Q_HEAT * 100.0

        Q_balance_out = (Q_loss_reactor+ Q_loss_tr1+ Q_loss_RTO+ Q_transport_loss_CO2+ Q_COOL+ N_EL_NET+ N_loss_em_total)
        Q_balance_error = Q_HEAT - Q_balance_out

        dTmin = rec["dTmin"]
        dT_T2_T6 = T1 - T6
        dT_T3_T5 = T3 - T5
        dT_reactor_hot = TR1 - T2_SET
        DP_REC_HOT = P2 - P3
        DP_REC_COLD = P5 - P6

        checks = {"P2<P1": P2 < P1,"P3<P2": P3 < P2,"P6<P5": P6 < P5,"P1=P6": abs(P1 - P6) < 1e-6,"TR2>T6": TR2 > T6,"dTmin=5K": abs(dTmin - DT_REC) < 1e-5,"T1>T3": T1 > T3,"T3>T5": T3 > T5,"T6>T5": T6 > T5,"T2_SET>T6": T2_SET > T6,"T1>T6": T1 > T6,"Q_RECUP>0": Q_RECUP > 0,"Q_COOL>0": Q_COOL > 0,"N_TURB_i>0": N_TURB_i > 0,"N_NET_i>0": N_NET_i > 0,"eta<100": eta_cycle_CO2 < 100,"balance<1kW": abs(Q_balance_error) < 1e3,"finite": np.all(np.isfinite([T2, T3, T5, T6, G_CO2,Q_RECUP, Q_COOL,N_TURB_i, N_MCOMP_i, N_NET_i,Q_balance_error]))}
        physical = all(checks.values())

        r.update({"P1_MPa": P1 / 1e6,"P2_MPa": P2 / 1e6,"P3_MPa": P3 / 1e6,"P4_MPa": P4 / 1e6,"P5_MPa": P5 / 1e6,"P6_MPa": P6 / 1e6,"T1_C": T1_C,"T2_before_transport_C": T2_SET_C,
            "T2_after_transport_C": T1_C,
            "deltaT_transport_CO2_K": T2_SET - T1,
            "T2_C": T2 - 273.15,
            "T3_C": T3 - 273.15,
            "T4_C": T4_C,"T5_C": T5 - 273.15,"T6_C": T6 - 273.15,"TR1_C": TR1_C,"TR2_C": TR2_C,"G_CO2_kg_s": G_CO2,"G_Na_kg_s": G_NA,"G_water_kg_s": G_WATER,
            "N_TURB_i_MW": N_TURB_i / 1e6,"N_MCOMP_i_MW": N_MCOMP_i / 1e6,"N_RCOMP_i_MW": N_RCOMP_i / 1e6,
            "N_NET_i_MW": N_NET_i / 1e6,
            "N_TURB_e_MW": N_TURB_e / 1e6,"N_MCOMP_e_MW": N_MCOMP_e / 1e6,"N_RCOMP_e_MW": N_RCOMP_e / 1e6,"N_TURB_el_MW": N_TURB_el / 1e6,"N_MCOMP_el_MW": N_MCOMP_el / 1e6,"N_RCOMP_el_MW": N_RCOMP_el / 1e6,
            "N_el_gross_MW": N_EL_GROSS / 1e6,
            "N_el_net_MW": N_EL_NET / 1e6,
            "Q_HEAT_MW": Q_HEAT / 1e6,
            "Q_reactor_MW": Q_reactor / 1e6,
            "Q_tr1_MW": Q_tr1 / 1e6,
            "Q_CO2_RTO_MW": Q_CO2_RTO / 1e6,
            "Q_HEAT_CALC_MW": Q_HEAT_CALC / 1e6,
            "Q_CO2_to_turbine_MW": Q_CO2_to_turbine / 1e6,
            "Q_transport_loss_CO2_MW": Q_transport_loss_CO2 / 1e6,
            "Q_RECUP_MW": Q_RECUP / 1e6,
            "Q_COOL_MW": Q_COOL / 1e6,
            "Q_loss_reactor_MW": Q_loss_reactor / 1e6,
            "Q_loss_tr1_MW": Q_loss_tr1 / 1e6,
            "Q_loss_RTO_MW": Q_loss_RTO / 1e6,
            "N_loss_turb_mech_MW": N_loss_turb_mech / 1e6,
            "N_loss_generator_MW": N_loss_generator / 1e6,
            "N_loss_turb_EM_MW": N_loss_turb_em / 1e6,
            "N_loss_mc_mech_MW": N_loss_mc_mech / 1e6,
            "N_loss_mc_drive_MW": N_loss_mc_drive / 1e6,
            "N_loss_comp_mech_MW": N_loss_comp_mech / 1e6,
            "N_loss_comp_drive_MW": N_loss_comp_drive / 1e6,
            "N_loss_comp_EM_MW": N_loss_comp_em / 1e6,
            "N_loss_pumps_MW": N_loss_pumps / 1e6,
            "N_loss_EM_total_MW": N_loss_em_total / 1e6,
            "N_Na_pump_kW": N_NA_PUMP / 1e3,"N_water_pump_kW": N_WATER_PUMP / 1e3,"N_SN_kW": N_SN / 1e3,"Q_balance_out_MW": Q_balance_out / 1e6,
            "Q_balance_error_MW": Q_balance_error / 1e6,"Q_balance_error_kW": Q_balance_error / 1e3,"eta_cycle_CO2_%": eta_cycle_CO2,"eta_plant_%": eta_plant,
            "eta_el_net_CO2_%": eta_el_net_CO2,"eta_el_net_plant_%": eta_el_net_plant,"dTmin_rec_K": dTmin,"dT_T2_T6_K": dT_T2_T6,"dT_T3_T5_K": dT_T3_T5,"dT_reactor_hot_K": dT_reactor_hot,
            "dT_reactor_cold_K": dT_reactor_cold,"DP_rec_hot_MPa": DP_REC_HOT / 1e6,"DP_rec_cold_MPa": DP_REC_COLD / 1e6,"energy_balance_MW": energy_balance / 1e6,"root_residual_max": root_residual_max,"physical_root": physical,"success": True,
            "message": "OK" if physical else "ROOT сошелся, но физическая проверка не пройдена.","T_hot_profile_C": [x - 273.15 for x in rec["T_hot_profile"]],"T_cold_profile_C": [x - 273.15 for x in rec["T_cold_profile"]],"dT_profile_K": rec["dT_profile"]})

        return r

    except Exception as e:
        r["message"] = f"{type(e).__name__}: {e}"
        return r


def worker(task):
    i, P1_calc, P2_calc = task
    try:
        return i, calc((P1_calc, P2_calc))
    except Exception as e:
        r = {
            "P1_MPa": P1_calc / 1e6,
            "P2_MPa": P2_calc / 1e6,
            "eta_T_%": ETA_T * 100,
            "eta_C_%": ETA_C * 100,
            "physical_root": False,
            "success": False,
            "message": f"Worker error: {type(e).__name__}: {e}"
        }
        return i, r


# ============================================================
# СОХРАНЕНИЕ В EXCEL
# ============================================================

def _write_results_sheet(ws, results):
    keys = list(results[0].keys())
    visible_keys = [k for k in keys if k not in ["T_hot_profile_C", "T_cold_profile_C", "dT_profile_K"]]
    ws.append(visible_keys)
    for r in results:
        row = []
        for k in visible_keys:
            value = r.get(k, np.nan)
            if k == "physical_root":
                value = bool(value)
            row.append(value)
        ws.append(row)
    results_ref = f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}"
    results_table = Table(displayName="ResultsTable", ref=results_ref)
    results_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(results_table)
    ws.auto_filter.ref = results_ref
    ws.freeze_panes = "A2"


def _write_best_block_sheet(ws, best):
    ws.title = "Best block"
    ws.append(["Параметр", "Турбина", "ОК", "РК", "Рекуператор", "Охладитель", "Насосы"])

    def add_row(name, *values):
        ws.append([name] + list(values))

    add_row("N_i, МВт",
            best.get("N_TURB_i_MW", ""),
            best.get("N_MCOMP_i_MW", ""),
            best.get("N_RCOMP_i_MW", ""),
            "", "", "")
    add_row("N_el, МВт",
            best.get("N_TURB_el_MW", ""),
            best.get("N_MCOMP_el_MW", ""),
            best.get("N_RCOMP_el_MW", ""),
            "", "", best.get("N_SN_kW", "") and best.get("N_SN_kW", 0) / 1000.0)
    add_row("Q_EM, МВт",
            best.get("N_loss_turb_EM_MW", ""),
            best.get("N_loss_comp_EM_MW", ""),
            "",
            "", "", best.get("N_loss_pumps_MW", ""))
    add_row("Q, МВт",
            "", "", "",
            best.get("Q_RECUP_MW", ""),
            best.get("Q_COOL_MW", ""),
            "")
    add_row("T вход, °C",
            best.get("T1_C", ""),
            best.get("T4_C", ""),
            "",
            best.get("T2_C", ""),
            best.get("T3_C", ""),
            "")
    add_row("T выход, °C",
            best.get("T2_C", ""),
            best.get("T5_C", ""),
            "",
            best.get("T6_C", ""),
            best.get("T4_C", ""),
            "")
    add_row("P вход, МПа",
            best.get("P1_MPa", ""),
            best.get("P4_MPa", ""),
            "",
            best.get("P2_MPa", ""),
            best.get("P3_MPa", ""),
            "")
    add_row("P выход, МПа",
            best.get("P2_MPa", ""),
            best.get("P5_MPa", ""),
            "",
            best.get("P3_MPa", ""),
            best.get("P4_MPa", ""),
            "")
    add_row("G, кг/с",
            best.get("G_CO2_kg_s", ""),
            best.get("G_CO2_kg_s", ""),
            "",
            best.get("G_CO2_kg_s", ""),
            best.get("G_water_kg_s", ""),
            "")
    add_row("", "", "", "", "", "", "")
    add_row("ЭНЕРГЕТИЧЕСКИЕ ПОКАЗАТЕЛИ", "", "", "", "", "", "")
    add_row("Q_HEAT, МВт", best.get("Q_HEAT_MW", ""), "", "", "", "", "")
    add_row("N_net_i, МВт", best.get("N_NET_i_MW", ""), "", "", "", "", "")
    add_row("N_el_gross, МВт", best.get("N_el_gross_MW", ""), "", "", "", "", "")
    add_row("N_el_net, МВт", best.get("N_el_net_MW", ""), "", "", "", "", "")
    add_row("N_SN, кВт", best.get("N_SN_kW", ""), "", "", "", "", "")
    add_row("eta_cycle_CO2, %", best.get("eta_cycle_CO2_%", ""), "", "", "", "", "")
    add_row("eta_el_net_CO2, %", best.get("eta_el_net_CO2_%", ""), "", "", "", "", "")
    add_row("eta_plant, %", best.get("eta_plant_%", ""), "", "", "", "", "")
    add_row("eta_el_net_plant, %", best.get("eta_el_net_plant_%", ""), "", "", "", "", "")

    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
        ws.column_dimensions[letter].width = min(width, 40)


def _write_balance_sheet(ws, best):
    ws.title = "Баланс"
    QH = best.get("Q_HEAT_MW", 1.0)

    def add_row(name, value_mw):
        if value_mw is None or value_mw == "":
            ws.append([name, "", ""])
        else:
            ws.append([name, value_mw, value_mw / QH * 100.0])

    ws.append(["Статья", "МВт", "% от Q_HEAT"])
    add_row("Q_loss_reactor", best.get("Q_loss_reactor_MW"))
    add_row("Q_loss_tr1 (Na)", best.get("Q_loss_tr1_MW"))
    add_row("Q_transport_loss_CO2", best.get("Q_transport_loss_CO2_MW"))

    # Сумма транспортных потерь
    q_tr1 = best.get("Q_loss_tr1_MW", 0) or 0
    q_tr2 = best.get("Q_transport_loss_CO2_MW", 0) or 0
    add_row("Транспортные потери (сумма)", q_tr1 + q_tr2)

    add_row("Q_loss_RTO", best.get("Q_loss_RTO_MW"))
    add_row("N_loss_turb_EM", best.get("N_loss_turb_EM_MW"))
    add_row("N_loss_comp_EM", best.get("N_loss_comp_EM_MW"))
    add_row("N_loss_pumps", best.get("N_loss_pumps_MW"))
    add_row("Q_COOL", best.get("Q_COOL_MW"))
    add_row("N_el_net", best.get("N_el_net_MW"))
    add_row("", "")
    add_row("Q_balance_out", best.get("Q_balance_out_MW"))
    add_row("Q_balance_error", best.get("Q_balance_error_MW"))
    add_row("Q_HEAT", QH)

    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
        ws.column_dimensions[letter].width = min(width, 40)


def save_excel(results):
    filename = "recuperative_CO2_DTmin5K_recuperator_losses.xlsx"
    wb = Workbook()

    ws_results = wb.active
    ws_results.title = "Results"
    _write_results_sheet(ws_results, results)

    physical = [r for r in results if r.get("physical_root", False)]
    if physical:
        best = max(physical, key=lambda r: r.get("eta_el_net_plant_%", -np.inf))
        ws_best = wb.create_sheet("Best block")
        _write_best_block_sheet(ws_best, best)
        ws_bal = wb.create_sheet("Баланс")
        _write_balance_sheet(ws_bal, best)

    for wsx in wb.worksheets:
        wsx.freeze_panes = "A2"
        for col in wsx.columns:
            letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            wsx.column_dimensions[letter].width = min(max_len + 2, 40)

    wb.save(filename)
    return filename


if __name__ == "__main__":
    mp.freeze_support()
    start = time.time()
    tasks = [(i, P1_calc, P2_calc) for i, (P1_calc, P2_calc) in enumerate(itertools.product(P1_RANGE, P2_RANGE))]
    nproc = min(14, mp.cpu_count(), len(tasks))
    results = [None] * len(tasks)
    print(f"Всего точек: {len(tasks)}")
    print(f"Процессов: {nproc}")
    print("=" * 75)
    with mp.Pool(nproc) as pool:
        for i, r in tqdm(
            pool.imap_unordered(worker, tasks, chunksize=1),
            total=len(tasks),
            desc="Расчёт",
            unit="точка"
        ):
            results[i] = r
    filename = save_excel(results)
    physical = [r for r in results if r.get("physical_root", False)]
    print(f"Всего решений: {len(results)}")
    print(f"Физических решений: {len(physical)}")
    if physical:
        best = max(physical, key=lambda r: r["eta_el_net_plant_%"])
        print()
        print("ЛУЧШАЯ ФИЗИЧЕСКАЯ ТОЧКА")
        print(f"P1 = {best['P1_MPa']:.3f} MPa")
        print(f"P2 = {best['P2_MPa']:.3f} MPa")
        print(f"T1 = {best['T1_C']:.3f} °C")
        print(f"T2 = {best['T2_C']:.3f} °C")
        print(f"T6 = {best['T6_C']:.3f} °C")
        print(f"G_CO2 = {best['G_CO2_kg_s']:.4f} kg/s")
        print(f"eta_cycle_CO2 = {best['eta_cycle_CO2_%']:.3f} %")
        print(f"eta_plant = {best['eta_plant_%']:.3f} %")
        print(f"eta_el_net_CO2 = {best['eta_el_net_CO2_%']:.3f} %")
        print(f"eta_el_net_plant = {best['eta_el_net_plant_%']:.3f} %")
        print(f"Q_balance_error = {best['Q_balance_error_MW']:.6e} MW")
    else:
        print()
        print("Физических решений нет.")
        print()
        print("Первые сообщения:")
        for r in results[:20]:
            print(f"P1={r.get('P1_MPa', np.nan):.1f}, P2={r.get('P2_MPa', np.nan):.1f}: {r.get('message', '')}")
    print()
    print(f"Excel: {filename}")
    print(f"Время: {(time.time() - start) / 60:.2f} мин")
    print(datetime.now())