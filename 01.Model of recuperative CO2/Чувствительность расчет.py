import numpy as np
import pandas as pd
from CoolProp.CoolProp import PropsSI as prop
from scipy.optimize import root
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
import time
from datetime import datetime


FLUID = "CO2"
TR1_C = 730.0
TR2_C = 550.0
T2_SET_C = 715.0
T1_C = 712.0
T4_C = 35.0
TR1 = TR1_C + 273.15
TR2 = TR2_C + 273.15
T2_SET = T2_SET_C + 273.15
T1 = T1_C + 273.15
T4 = T4_C + 273.15
Q_HEAT = 2.750e6
KPDcomp = 0.80
KPDturb = 0.90
KPD_reactor = 0.99
KPD_RTO = 0.99
ETA_GENERATOR = 0.98
eta_turb_mech = 0.99
eta_generator = 0.99
eta_mc_mech = 0.97
eta_rc_mech = 0.97
eta_drive = 0.98

eta_tr1 = 0.99
eta_tr2 = 0.99

P1_FIXED_MPa = 28.0
P2_FIXED_MPa = 12.0
P1_FIXED = P1_FIXED_MPa * 1e6
P2_FIXED = P2_FIXED_MPa * 1e6
N_REC_SECTIONS = 20

CP_NA = 1270.0
RHO_NA = 850.0
RHO_WATER = 1000.0
CP_WATER = 4180.0
P_NA = 1.0e6
DP_NA_FRAC = 0.05
P_WATER = 0.50e6
DP_WATER_FRAC = 0.10
ETA_PUMP = 0.80
ETA_MOTOR = 0.95
DT_WATER = 10.0

Q_reactor = Q_HEAT * KPD_reactor
Q_tr1 = Q_reactor * eta_tr1
Q_CO2_RTO = Q_tr1 * KPD_RTO
Q_loss_reactor = Q_HEAT - Q_reactor
Q_loss_tr1 = Q_reactor - Q_tr1
Q_loss_RTO = Q_tr1 - Q_CO2_RTO
G_NA = Q_reactor / (CP_NA * (TR1_C - TR2_C))

ETA_T_BASE = KPDturb
ETA_C_BASE = KPDcomp
DP_REC_HOT_BASE = 0.03
DP_REC_COLD_BASE = 0.02
DT_REC_BASE_K = 5.0

ETA_T_MIN = 0.80
ETA_T_MAX = 0.95
ETA_T_STEP = 0.01

ETA_C_MIN = 0.70
ETA_C_MAX = 0.90
ETA_C_STEP = 0.01

DP_HOT_MIN = 0.00
DP_HOT_MAX = 0.15
DP_HOT_STEP = 0.01

DP_COLD_MIN = 0.00
DP_COLD_MAX = 0.15
DP_COLD_STEP = 0.01

DT_REC_MIN_K = 3.0
DT_REC_MAX_K = 15.0
DT_REC_STEP_K = 0.1

ETA_T_VALUES = np.round(
    np.arange(ETA_T_MIN, ETA_T_MAX + ETA_T_STEP / 2, ETA_T_STEP), 2
)

ETA_C_VALUES = np.round(
    np.arange(ETA_C_MIN, ETA_C_MAX + ETA_C_STEP / 2, ETA_C_STEP), 2
)

DP_HOT_VALUES = np.round(
    np.arange(DP_HOT_MIN, DP_HOT_MAX + DP_HOT_STEP / 2, DP_HOT_STEP), 2
)

DP_COLD_VALUES = np.round(
    np.arange(DP_COLD_MIN, DP_COLD_MAX + DP_COLD_STEP / 2, DP_COLD_STEP), 2
)

DT_REC_VALUES = np.round(
    np.arange(DT_REC_MIN_K, DT_REC_MAX_K + DT_REC_STEP_K / 2, DT_REC_STEP_K), 1
)

# ============================================================
# ТЕРМОДИНАМИЧЕСКИЕ ФУНКЦИИ
# ============================================================

def H_TP(T, P):
    return prop("H", "T", T, "P", P, FLUID)

def S_TP(T, P):
    return prop("S", "T", T, "P", P, FLUID)

def H_PS(P, S):
    return prop("H", "P", P, "S", S, FLUID)

def T_PH(P, H):
    return prop("T", "P", P, "H", H, FLUID)

# ============================================================
# ДАВЛЕНИЯ
# ============================================================

def calculate_pressures(P1, P2, dp_hot, dp_cold):
    if P2 >= P1:
        raise ValueError("P2 >= P1")

    P3 = P2 * (1.0 - dp_hot)
    P4 = P3
    P6 = P1
    P5 = P6 / (1.0 - dp_cold)

    return {
        "P1": P1,
        "P2": P2,
        "P3": P3,
        "P4": P4,
        "P5": P5,
        "P6": P6
    }

# ============================================================
# ПРОФИЛЬ РЕКУПЕРАТОРА
# ============================================================

def recuperator_profile(P2, P3, P5, P6, h2, h5, T6):
    h6 = H_TP(T6, P6)
    dh_cold = h6 - h5

    if dh_cold <= 0:
        raise ValueError("h6-h5 <= 0: холодный поток не получает тепло")

    h3 = h2 - dh_cold
    T3 = T_PH(P3, h3)

    frac = np.linspace(0.0, 1.0, N_REC_SECTIONS + 1)

    H_hot = h2 - frac * (h2 - h3)
    H_cold = h6 - frac * (h6 - h5)

    T_hot = np.array([T_PH(P2, h) for h in H_hot])
    T_cold = np.array([T_PH(P6, h) for h in H_cold])

    dT_profile = T_hot - T_cold
    dTmin = float(np.min(dT_profile))
    idx_min = int(np.argmin(dT_profile))

    return {
        "h6": h6,
        "h3": h3,
        "T3": T3,
        "Q_rec_specific": dh_cold,
        "T_hot_profile": T_hot.tolist(),
        "T_cold_profile": T_cold.tolist(),
        "dT_profile": dT_profile.tolist(),
        "dTmin": dTmin,
        "idx_min": idx_min
    }

# ============================================================
# СТРУКТУРА РЕЗУЛЬТАТА
# ============================================================

OUTPUT_COLUMNS = [
    "eta_T_%",
    "eta_C_%",
    "DP_rec_hot_%",
    "DP_rec_cold_%",
    "dTmin_target_K",
    "P1_MPa",
    "P2_MPa",
    "P3_MPa",
    "P4_MPa",
    "P5_MPa",
    "P6_MPa",
    "T1_C",
    "T2_before_transport_C",
    "T2_after_transport_C",
    "deltaT_transport_CO2_K",
    "T2_C",
    "T3_C",
    "T4_C",
    "T5_C",
    "T6_C",
    "TR1_C",
    "TR2_C",
    "G_CO2_kg_s",
    "G_Na_kg_s",
    "G_water_kg_s",
    "N_TURB_i_MW",
    "N_MCOMP_i_MW",
    "N_RCOMP_i_MW",
    "N_NET_i_MW",
    "N_TURB_e_MW",
    "N_MCOMP_e_MW",
    "N_RCOMP_e_MW",
    "N_el_gross_MW",
    "N_el_net_MW",
    "Q_HEAT_MW",
    "Q_reactor_MW",
    "Q_tr1_MW",
    "Q_CO2_RTO_MW",
    "Q_HEAT_CALC_MW",
    "Q_CO2_to_turbine_MW",
    "Q_transport_loss_CO2_MW",
    "Q_RECUP_MW",
    "Q_COOL_MW",
    "Q_loss_reactor_MW",
    "Q_loss_tr1_MW",
    "Q_loss_RTO_MW",
    "N_loss_turb_mech_MW",
    "N_loss_generator_MW",
    "N_loss_turb_EM_MW",
    "N_loss_mc_mech_MW",
    "N_loss_mc_drive_MW",
    "N_loss_comp_mech_MW",
    "N_loss_comp_drive_MW",
    "N_loss_comp_EM_MW",
    "N_loss_pumps_MW",
    "N_loss_EM_total_MW",
    "N_Na_pump_kW",
    "N_water_pump_kW",
    "N_SN_kW",
    "Q_balance_out_MW",
    "Q_balance_error_MW",
    "Q_balance_error_kW",
    "eta_cycle_CO2_%",
    "eta_plant_%",
    "eta_el_net_CO2_%",
    "eta_el_net_plant_%",
    "dTmin_rec_K",
    "dT_T2_T6_K",
    "dT_T3_T5_K",
    "dT_reactor_hot_K",
    "dT_reactor_cold_K",
    "DP_rec_hot_MPa",
    "DP_rec_cold_MPa",
    "energy_balance_MW",
    "root_residual_max",
    "physical_root",
    "success",
    "message"
]

def empty_result(eta_t, eta_c, dp_hot, dp_cold, dt_rec, message=""):
    r = {key: np.nan for key in OUTPUT_COLUMNS}

    r["eta_T_%"] = eta_t * 100.0
    r["eta_C_%"] = eta_c * 100.0
    r["DP_rec_hot_%"] = dp_hot * 100.0
    r["DP_rec_cold_%"] = dp_cold * 100.0
    r["dTmin_target_K"] = dt_rec

    r["P1_MPa"] = P1_FIXED_MPa
    r["P2_MPa"] = P2_FIXED_MPa

    r["T1_C"] = T1_C
    r["T2_before_transport_C"] = T2_SET_C
    r["T2_after_transport_C"] = T1_C
    r["T4_C"] = T4_C
    r["TR1_C"] = TR1_C
    r["TR2_C"] = TR2_C

    r["Q_HEAT_MW"] = Q_HEAT / 1e6
    r["Q_reactor_MW"] = Q_reactor / 1e6
    r["Q_tr1_MW"] = Q_tr1 / 1e6
    r["Q_CO2_RTO_MW"] = Q_CO2_RTO / 1e6
    r["Q_loss_reactor_MW"] = Q_loss_reactor / 1e6
    r["Q_loss_tr1_MW"] = Q_loss_tr1 / 1e6
    r["Q_loss_RTO_MW"] = Q_loss_RTO / 1e6

    r["physical_root"] = False
    r["success"] = False
    r["message"] = message

    return r

# ============================================================
# ОСНОВНОЙ РАСЧЁТ ОДНОЙ ТОЧКИ
# ============================================================

def calculate_case(P1, P2, eta_t, eta_c, dp_hot, dp_cold, dt_rec):

    r = empty_result(eta_t, eta_c, dp_hot, dp_cold, dt_rec)

    try:
        pressures = calculate_pressures(P1, P2, dp_hot, dp_cold)

        P1 = pressures["P1"]
        P2 = pressures["P2"]
        P3 = pressures["P3"]
        P4 = pressures["P4"]
        P5 = pressures["P5"]
        P6 = pressures["P6"]

        # ----------------------------------------------------
        # Турбина
        # ----------------------------------------------------

        h1 = H_TP(T1, P1)
        s1 = S_TP(T1, P1)

        h2s = H_PS(P2, s1)
        h2 = h1 - eta_t * (h1 - h2s)

        T2 = T_PH(P2, h2)
        w_turb = h1 - h2

        # ----------------------------------------------------
        # Главный компрессор
        # ----------------------------------------------------

        h4 = H_TP(T4, P4)
        s4 = S_TP(T4, P4)

        h5s = H_PS(P5, s4)
        h5 = h4 + (h5s - h4) / eta_c

        T5 = T_PH(P5, h5)
        w_comp = h5 - h4

        # ----------------------------------------------------
        # Транспорт СО2
        # ----------------------------------------------------

        h_before_transport = H_TP(T2_SET, P1)
        h_after_transport = H_TP(T1, P1)

        # ----------------------------------------------------
        # Система уравнений
        # ----------------------------------------------------

        def equations(x):
            G = float(x[0])
            T6_local = float(x[1])

            if G <= 0:
                return np.array([1e6, 1e6])

            try:
                h6_local = H_TP(T6_local, P6)

                eq1 = (
                    G * (h_before_transport - h6_local)
                    - Q_CO2_RTO
                )

                rec_local = recuperator_profile(
                    P2, P3, P5, P6,
                    h2, h5, T6_local
                )

                eq2 = rec_local["dTmin"] - dt_rec

                return np.array([
                    eq1 / 1e6,
                    eq2
                ])

            except Exception:
                return np.array([1e6, 1e6])

        # ----------------------------------------------------
        # Начальное приближение
        # ----------------------------------------------------

        G0 = Q_CO2_RTO / (
            h_before_transport
            - H_TP(TR2 - 5.0, P6)
        )

        if not np.isfinite(G0) or G0 <= 0:
            G0 = 20.0

        x0 = np.array([
            G0,
            TR2 - 5.0
        ])

        sol = root(
            equations,
            x0=x0,
            method="hybr",
            tol=1e-8
        )

        if not sol.success:
            raise ValueError(
                "ROOT не сошелся: " + str(sol.message)
            )

        G_CO2 = float(sol.x[0])
        T6 = float(sol.x[1])

        residual = equations([G_CO2, T6])
        root_residual_max = float(
            np.max(np.abs(residual))
        )

        if root_residual_max > 1e-5:
            raise ValueError(
                f"ROOT дал слишком большую невязку: "
                f"{root_residual_max:.3e}"
            )

        # ----------------------------------------------------
        # Проверка температуры реактора
        # ----------------------------------------------------

        dT_reactor_cold = TR2 - T6

        if dT_reactor_cold <= 0:
            raise ValueError(
                f"TR2 > T6 не выполняется: "
                f"TR2={TR2_C:.3f} C, "
                f"T6={T6 - 273.15:.3f} C"
            )

        # ----------------------------------------------------
        # Рекуператор
        # ----------------------------------------------------

        rec = recuperator_profile(
            P2, P3, P5, P6,
            h2, h5, T6
        )

        h6 = rec["h6"]
        h3 = rec["h3"]
        T3 = rec["T3"]

        Q_RECUP = G_CO2 * rec["Q_rec_specific"]

        Q_COOL = G_CO2 * (h3 - h4)

        Q_HEAT_CALC = G_CO2 * (
            h_before_transport - h6
        )

        # ----------------------------------------------------
        # Потери транспорта СО2
        # ----------------------------------------------------

        Q_transport_loss_CO2 = G_CO2 * (
            h_before_transport - h_after_transport
        )

        Q_CO2_to_turbine = (
            Q_CO2_RTO
            - Q_transport_loss_CO2
        )

        # ----------------------------------------------------
        # Механическая мощность
        # ----------------------------------------------------

        N_TURB_i = G_CO2 * w_turb
        N_MCOMP_i = G_CO2 * w_comp
        N_RCOMP_i = 0.0

        N_NET_i = (
            N_TURB_i
            - N_MCOMP_i
            - N_RCOMP_i
        )

        energy_balance = (
            Q_CO2_to_turbine
            - Q_COOL
            - N_NET_i
        )

        # ----------------------------------------------------
        # Насосы
        # ----------------------------------------------------

        DP_NA = P_NA * DP_NA_FRAC

        N_NA_PUMP = (
            G_NA
            / RHO_NA
            * DP_NA
            / ETA_PUMP
            / ETA_MOTOR
        )

        G_WATER = (
            Q_COOL
            / (CP_WATER * DT_WATER)
        )

        DP_WATER = P_WATER * DP_WATER_FRAC

        N_WATER_PUMP = (
            G_WATER
            / RHO_WATER
            * DP_WATER
            / ETA_PUMP
            / ETA_MOTOR
        )

        N_SN = (
            N_NA_PUMP
            + N_WATER_PUMP
        )

        # ----------------------------------------------------
        # Электрическая мощность
        # ----------------------------------------------------

        N_TURB_e = (
            N_TURB_i
            * eta_turb_mech
        )

        N_TURB_el = (
            N_TURB_e
            * eta_generator
        )

        N_MCOMP_e = (
            N_MCOMP_i
            / eta_mc_mech
        )

        N_MCOMP_el = (
            N_MCOMP_e
            / eta_drive
        )

        N_RCOMP_e = 0.0
        N_RCOMP_el = 0.0

        N_EL_GROSS = (
            N_TURB_el
            - N_MCOMP_el
            - N_RCOMP_el
        )

        N_EL_NET = (
            N_EL_GROSS
            - N_SN
        )

        # ----------------------------------------------------
        # Электромеханические потери
        # ----------------------------------------------------

        N_loss_turb_mech = (
            N_TURB_i
            - N_TURB_e
        )

        N_loss_generator = (
            N_TURB_e
            - N_TURB_el
        )

        N_loss_turb_em = (
            N_loss_turb_mech
            + N_loss_generator
        )

        N_loss_mc_mech = (
            N_MCOMP_e
            - N_MCOMP_i
        )

        N_loss_mc_drive = (
            N_MCOMP_el
            - N_MCOMP_e
        )

        N_loss_rc_mech = 0.0
        N_loss_rc_drive = 0.0

        N_loss_comp_mech = (
            N_loss_mc_mech
            + N_loss_rc_mech
        )

        N_loss_comp_drive = (
            N_loss_mc_drive
            + N_loss_rc_drive
        )

        N_loss_comp_em = (
            N_loss_comp_mech
            + N_loss_comp_drive
        )

        N_loss_pumps = N_SN

        N_loss_em_total = (
            N_loss_turb_em
            + N_loss_comp_em
            + N_loss_pumps
        )

        # ----------------------------------------------------
        # КПД
        # ----------------------------------------------------

        eta_cycle_CO2 = (
            N_NET_i
            / Q_CO2_RTO
            * 100.0
        )

        eta_el_net_CO2 = (
            N_EL_NET
            / Q_CO2_RTO
            * 100.0
        )

        eta_plant = (
            N_NET_i
            / Q_HEAT
            * 100.0
        )

        eta_el_net_plant = (
            N_EL_NET
            / Q_HEAT
            * 100.0
        )

        # ----------------------------------------------------
        # Энергетический баланс
        # ----------------------------------------------------

        Q_balance_out = (
            Q_loss_reactor
            + Q_loss_tr1
            + Q_loss_RTO
            + Q_transport_loss_CO2
            + Q_COOL
            + N_EL_NET
            + N_loss_em_total
        )

        Q_balance_error = (
            Q_HEAT
            - Q_balance_out
        )

        # ----------------------------------------------------
        # Дополнительные показатели
        # ----------------------------------------------------

        dTmin = rec["dTmin"]

        dT_T2_T6 = T1 - T6
        dT_T3_T5 = T3 - T5

        dT_reactor_hot = (
            TR1 - T2_SET
        )

        DP_REC_HOT = P2 - P3
        DP_REC_COLD = P5 - P6

        # ----------------------------------------------------
        # Физические проверки
        # ----------------------------------------------------

        checks = {
            "P2<P1": P2 < P1,
            "P3<P2": P3 < P2,
            "P6<P5": P6 < P5,
            "P1=P6": abs(P1 - P6) < 1e-6,
            "TR2>T6": TR2 > T6,
            "dTmin": abs(dTmin - dt_rec) < 1e-5,
            "T1>T3": T1 > T3,
            "T3>T5": T3 > T5,
            "T6>T5": T6 > T5,
            "T2_SET>T6": T2_SET > T6,
            "T1>T6": T1 > T6,
            "Q_RECUP>0": Q_RECUP > 0,
            "Q_COOL>0": Q_COOL > 0,
            "N_TURB_i>0": N_TURB_i > 0,
            "N_NET_i>0": N_NET_i > 0,
            "eta<100": eta_cycle_CO2 < 100,
            "balance<1kW": abs(Q_balance_error) < 1e3,
            "finite": np.all(np.isfinite([
                T2,
                T3,
                T5,
                T6,
                G_CO2,
                Q_RECUP,
                Q_COOL,
                N_TURB_i,
                N_MCOMP_i,
                N_NET_i,
                Q_balance_error
            ]))
        }

        physical = all(checks.values())

        # ----------------------------------------------------
        # Результат
        # ----------------------------------------------------

        r.update({
            "P1_MPa": P1 / 1e6,
            "P2_MPa": P2 / 1e6,
            "P3_MPa": P3 / 1e6,
            "P4_MPa": P4 / 1e6,
            "P5_MPa": P5 / 1e6,
            "P6_MPa": P6 / 1e6,

            "T1_C": T1_C,
            "T2_before_transport_C": T2_SET_C,
            "T2_after_transport_C": T1_C,
            "deltaT_transport_CO2_K": T2_SET - T1,

            "T2_C": T2 - 273.15,
            "T3_C": T3 - 273.15,
            "T4_C": T4_C,
            "T5_C": T5 - 273.15,
            "T6_C": T6 - 273.15,

            "TR1_C": TR1_C,
            "TR2_C": TR2_C,

            "G_CO2_kg_s": G_CO2,
            "G_Na_kg_s": G_NA,
            "G_water_kg_s": G_WATER,

            "N_TURB_i_MW": N_TURB_i / 1e6,
            "N_MCOMP_i_MW": N_MCOMP_i / 1e6,
            "N_RCOMP_i_MW": N_RCOMP_i / 1e6,
            "N_NET_i_MW": N_NET_i / 1e6,

            "N_TURB_e_MW": N_TURB_e / 1e6,
            "N_MCOMP_e_MW": N_MCOMP_e / 1e6,
            "N_RCOMP_e_MW": N_RCOMP_e / 1e6,

            "N_el_gross_MW": N_EL_GROSS / 1e6,
            "N_el_net_MW": N_EL_NET / 1e6,

            "Q_HEAT_MW": Q_HEAT / 1e6,
            "Q_reactor_MW": Q_reactor / 1e6,
            "Q_tr1_MW": Q_tr1 / 1e6,
            "Q_CO2_RTO_MW": Q_CO2_RTO / 1e6,

            "Q_HEAT_CALC_MW": Q_HEAT_CALC / 1e6,
            "Q_CO2_to_turbine_MW": Q_CO2_to_turbine / 1e6,
            "Q_transport_loss_CO2_MW": Q_transport_loss_CO2 / 1e6,

            "Q_RECUP_MW": Q_RECUP / 1e6,
            "Q_COOL_MW": Q_COOL / 1e6,

            "Q_loss_reactor_MW": Q_loss_reactor / 1e6,
            "Q_loss_tr1_MW": Q_loss_tr1 / 1e6,
            "Q_loss_RTO_MW": Q_loss_RTO / 1e6,

            "N_loss_turb_mech_MW": N_loss_turb_mech / 1e6,
            "N_loss_generator_MW": N_loss_generator / 1e6,
            "N_loss_turb_EM_MW": N_loss_turb_em / 1e6,

            "N_loss_mc_mech_MW": N_loss_mc_mech / 1e6,
            "N_loss_mc_drive_MW": N_loss_mc_drive / 1e6,

            "N_loss_comp_mech_MW": N_loss_comp_mech / 1e6,
            "N_loss_comp_drive_MW": N_loss_comp_drive / 1e6,
            "N_loss_comp_EM_MW": N_loss_comp_em / 1e6,

            "N_loss_pumps_MW": N_loss_pumps / 1e6,
            "N_loss_EM_total_MW": N_loss_em_total / 1e6,

            "N_Na_pump_kW": N_NA_PUMP / 1e3,
            "N_water_pump_kW": N_WATER_PUMP / 1e3,
            "N_SN_kW": N_SN / 1e3,

            "Q_balance_out_MW": Q_balance_out / 1e6,
            "Q_balance_error_MW": Q_balance_error / 1e6,
            "Q_balance_error_kW": Q_balance_error / 1e3,

            "eta_cycle_CO2_%": eta_cycle_CO2,
            "eta_plant_%": eta_plant,
            "eta_el_net_CO2_%": eta_el_net_CO2,
            "eta_el_net_plant_%": eta_el_net_plant,

            "dTmin_rec_K": dTmin,
            "dT_T2_T6_K": dT_T2_T6,
            "dT_T3_T5_K": dT_T3_T5,
            "dT_reactor_hot_K": dT_reactor_hot,
            "dT_reactor_cold_K": dT_reactor_cold,

            "DP_rec_hot_MPa": DP_REC_HOT / 1e6,
            "DP_rec_cold_MPa": DP_REC_COLD / 1e6,

            "energy_balance_MW": energy_balance / 1e6,
            "root_residual_max": root_residual_max,

            "physical_root": physical,
            "success": True,
            "message": (
                "OK"
                if physical
                else "ROOT сошелся, но физическая проверка не пройдена."
            )
        })

        return r

    except Exception as e:
        r["message"] = f"{type(e).__name__}: {e}"
        return r

# ============================================================
# DATAFRAME
# ============================================================

def results_to_dataframe(results):
    df = pd.DataFrame(results)

    for column in OUTPUT_COLUMNS:
        if column not in df.columns:
            df[column] = np.nan

    df = df[OUTPUT_COLUMNS]

    return df

# ============================================================
# EXCEL
# ============================================================

def write_dataframe_to_sheet(wb, df, sheet_name, table_name):
    ws = wb.create_sheet(sheet_name)

    ws.append(list(df.columns))

    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))

    if ws.max_row >= 2 and ws.max_column >= 1:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

        table = Table(
            displayName=table_name,
            ref=ref
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        ws.add_table(table)
        ws.auto_filter.ref = ref

    ws.freeze_panes = "A2"

    for column in ws.columns:
        letter = get_column_letter(column[0].column)

        max_length = 0

        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        ws.column_dimensions[letter].width = min(
            max(max_length + 2, 10),
            32
        )

# ============================================================
# SUMMARY
# ============================================================

def make_summary(df_eta, df_dp, df_dt):

    rows = []

    def add_summary(name, df):
        total = len(df)
        success = int(df["success"].sum())
        physical = int(df["physical_root"].sum())

        rows.append([
            name,
            total,
            success,
            physical
        ])

        if physical > 0:
            physical_df = df[df["physical_root"] == True]

            eta_max = physical_df["eta_el_net_plant_%"].max()
            eta_min = physical_df["eta_el_net_plant_%"].min()

            rows.append([
                f"{name}: диапазон η_el_net_plant",
                "",
                "",
                f"{eta_min:.6f}–{eta_max:.6f} %"
            ])

    add_summary(
        "ηT × ηC",
        df_eta
    )

    add_summary(
        "ΔP рекуператора",
        df_dp
    )

    add_summary(
        "ΔTmin",
        df_dt
    )

    return rows

# ============================================================
# ЗАПИСЬ SUMMARY
# ============================================================

def write_summary(wb, df_eta, df_dp, df_dt):

    ws = wb.create_sheet("Summary")

    ws.append([
        "Параметр",
        "Значение"
    ])

    summary_params = [
        ("Рабочее тело", FLUID),
        ("P1, МПа", P1_FIXED_MPa),
        ("P2, МПа", P2_FIXED_MPa),
        ("Q_HEAT, МВт", Q_HEAT / 1e6),
        ("TR1, °C", TR1_C),
        ("TR2, °C", TR2_C),
        ("T2 задаваемая, °C", T2_SET_C),
        ("T1, °C", T1_C),
        ("T4, °C", T4_C),
        ("Базовый ηT, %", ETA_T_BASE * 100),
        ("Базовый ηC, %", ETA_C_BASE * 100),
        ("Базовые ΔP горячей стороны, %", DP_REC_HOT_BASE * 100),
        ("Базовые ΔP холодной стороны, %", DP_REC_COLD_BASE * 100),
        ("Базовый ΔTmin, K", DT_REC_BASE_K),
        ("Диапазон ηT, %", f"{ETA_T_MIN * 100:.0f}–{ETA_T_MAX * 100:.0f}"),
        ("Диапазон ηC, %", f"{ETA_C_MIN * 100:.0f}–{ETA_C_MAX * 100:.0f}"),
        ("Диапазон ΔP горячей стороны, %", f"{DP_HOT_MIN * 100:.0f}–{DP_HOT_MAX * 100:.0f}"),
        ("Диапазон ΔP холодной стороны, %", f"{DP_COLD_MIN * 100:.0f}–{DP_COLD_MAX * 100:.0f}"),
        ("Диапазон ΔTmin, K", f"{DT_REC_MIN_K:.1f}–{DT_REC_MAX_K:.1f}"),
        ("Шаг ΔTmin, K", DT_REC_STEP_K),
        ("Число секций рекуператора", N_REC_SECTIONS),
        ("Количество точек ηT × ηC", len(df_eta)),
        ("Количество точек ΔP", len(df_dp)),
        ("Количество точек ΔTmin", len(df_dt))
    ]

    for name, value in summary_params:
        ws.append([
            name,
            value
        ])

    ws.append([])
    ws.append([
        "Анализ",
        "Всего точек",
        "ROOT сошелся",
        "Физических решений"
    ])

    summary_rows = make_summary(
        df_eta,
        df_dp,
        df_dt
    )

    for row in summary_rows:
        ws.append(row)

    ws.freeze_panes = "A2"

    for column in ws.columns:
        letter = get_column_letter(column[0].column)

        max_length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in column
        )

        ws.column_dimensions[letter].width = min(
            max(max_length + 2, 12),
            45
        )

# ============================================================
# ОСНОВНОЙ ЗАПУСК
# ============================================================

if __name__ == "__main__":

    start_time = time.time()

    print("ЧУВСТВИТЕЛЬНОСТЬ РЕКУПЕРАТИВНОГО CO2-ЦИКЛА")
    print(f"P1 = {P1_FIXED_MPa:.1f} МПа")
    print(f"P2 = {P2_FIXED_MPa:.1f} МПа")
    print(f"ηT base = {ETA_T_BASE * 100:.1f} %")
    print(f"ηC base = {ETA_C_BASE * 100:.1f} %")
    print(f"ΔP hot base = {DP_REC_HOT_BASE * 100:.1f} %")
    print(f"ΔP cold base = {DP_REC_COLD_BASE * 100:.1f} %")
    print(f"ΔTmin base = {DT_REC_BASE_K:.1f} K")
    print()

    # ========================================================
    # АНАЛИЗ 1: ηT × ηC
    # ========================================================

    print("=" * 70)
    print("АНАЛИЗ 1: КПД ТУРБИНЫ И КОМПРЕССОРА")
    print("=" * 70)

    results_eta = []

    total_eta = (
        len(ETA_T_VALUES)
        * len(ETA_C_VALUES)
    )

    counter = 0

    for eta_t in ETA_T_VALUES:

        for eta_c in ETA_C_VALUES:

            counter += 1

            print(
                f"\rηT/ηC: {counter}/{total_eta} "
                f"({counter / total_eta * 100:.1f}%)",
                end=""
            )

            r = calculate_case(
                P1_FIXED,
                P2_FIXED,
                eta_t,
                eta_c,
                DP_REC_HOT_BASE,
                DP_REC_COLD_BASE,
                DT_REC_BASE_K
            )

            results_eta.append(r)

    print()

    df_eta = results_to_dataframe(
        results_eta
    )

    # ========================================================
    # АНАЛИЗ 2: ПОТЕРИ ДАВЛЕНИЯ
    # ========================================================

    print()
    print("=" * 70)
    print("АНАЛИЗ 2: ПОТЕРИ ДАВЛЕНИЯ РЕКУПЕРАТОРА")
    print("=" * 70)

    results_dp = []

    total_dp = (
        len(DP_HOT_VALUES)
        * len(DP_COLD_VALUES)
    )

    counter = 0

    for dp_hot in DP_HOT_VALUES:

        for dp_cold in DP_COLD_VALUES:

            counter += 1

            print(
                f"\rΔP: {counter}/{total_dp} "
                f"({counter / total_dp * 100:.1f}%)",
                end=""
            )

            r = calculate_case(
                P1_FIXED,
                P2_FIXED,
                ETA_T_BASE,
                ETA_C_BASE,
                dp_hot,
                dp_cold,
                DT_REC_BASE_K
            )

            results_dp.append(r)

    print()

    df_dp = results_to_dataframe(
        results_dp
    )

    # ========================================================
    # АНАЛИЗ 3: ΔTmin
    # ========================================================

    print()
    print("=" * 70)
    print("АНАЛИЗ 3: МИНИМАЛЬНЫЙ ТЕМПЕРАТУРНЫЙ НАПОР")
    print("=" * 70)

    results_dt = []

    total_dt = len(
        DT_REC_VALUES
    )

    counter = 0

    for dt_rec in DT_REC_VALUES:

        counter += 1

        print(
            f"\rΔTmin: {counter}/{total_dt} "
            f"({counter / total_dt * 100:.1f}%)",
            end=""
        )

        r = calculate_case(
            P1_FIXED,
            P2_FIXED,
            ETA_T_BASE,
            ETA_C_BASE,
            DP_REC_HOT_BASE,
            DP_REC_COLD_BASE,
            dt_rec
        )

        results_dt.append(r)

    print()

    df_dt = results_to_dataframe(
        results_dt
    )

    # ========================================================
    # EXCEL
    # ========================================================

    filename = (
        "sensitivity_CO2_recuperative_cycle.xlsx"
    )

    wb = Workbook()

    default_ws = wb.active
    wb.remove(default_ws)

    write_dataframe_to_sheet(
        wb,
        df_eta,
        "Results_eta",
        "EtaSensitivity"
    )

    write_dataframe_to_sheet(
        wb,
        df_dp,
        "Results_pressure",
        "PressureSensitivity"
    )

    write_dataframe_to_sheet(
        wb,
        df_dt,
        "Results_dT",
        "DTSensitivity"
    )

    write_summary(
        wb,
        df_eta,
        df_dp,
        df_dt
    )

    wb.save(filename)

    # ========================================================
    # ИТОГ
    # ========================================================

    print()
    print("=" * 70)
    print("РАСЧЁТ ЗАВЕРШЁН")
    print("=" * 70)

    print(
        f"ηT × ηC: {len(df_eta)} точек"
    )

    print(
        f"ΔP hot × ΔP cold: {len(df_dp)} точек"
    )

    print(
        f"ΔTmin: {len(df_dt)} точек"
    )

    print(
        f"ROOT сошелся: "
        f"{int(df_eta['success'].sum())}/"
        f"{len(df_eta)}; "
        f"{int(df_dp['success'].sum())}/"
        f"{len(df_dp)}; "
        f"{int(df_dt['success'].sum())}/"
        f"{len(df_dt)}"
    )

    print(
        f"Физических решений: "
        f"{int(df_eta['physical_root'].sum())}/"
        f"{len(df_eta)}; "
        f"{int(df_dp['physical_root'].sum())}/"
        f"{len(df_dp)}; "
        f"{int(df_dt['physical_root'].sum())}/"
        f"{len(df_dt)}"
    )

    print()
    print(f"Excel: {filename}")
    print(
        f"Время расчёта: "
        f"{(time.time() - start_time) / 60:.2f} мин"
    )
    print(datetime.now())